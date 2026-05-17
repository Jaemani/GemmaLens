SAMPLE_TEXT = (
    "Although previous studies have suggested a correlation between sleep deprivation "
    "and reduced cognitive performance, the extent to which these findings generalize "
    "across real-world learning environments remains unclear. To address this gap, "
    "we analyze longitudinal study logs collected from undergraduate students over a "
    "six-week period."
)


def test_health_and_model_status(client):
    health = client.get("/health")
    assert health.status_code == 200
    assert health.json() == {"status": "ok"}

    status = client.get("/models/status")
    assert status.status_code == 200
    body = status.json()
    assert body["mock_fallback"] is True
    assert "provider" in body
    assert "mlx_model_available" in body


def test_profile_language_settings(client):
    profile = client.get("/profile")
    assert profile.status_code == 200
    assert profile.json()["learning_language"] == "English"

    updated = client.patch(
        "/profile",
        json={
            "support_language": "Japanese",
            "learning_language": "Spanish",
            "target_level": "B2",
            "onboarding_completed": True,
        },
    )
    assert updated.status_code == 200
    body = updated.json()
    assert body["support_language"] == "Japanese"
    assert body["learning_language"] == "Spanish"
    assert body["target_level"] == "B2"
    assert body["onboarding_completed"] is True


def test_analysis_uses_profile_level_when_model_level_unknown(client):
    updated = client.patch("/profile", json={"target_level": "C2"})
    assert updated.status_code == 200
    created = client.post(
        "/documents",
        json={"title": "Short ML note", "content": "We introduce a new language representation model called BERT.", "source_type": "text"},
    )
    assert created.status_code == 200

    analyzed = client.post(f"/documents/{created.json()['id']}/analyze")
    assert analyzed.status_code == 200
    assert analyzed.json()["difficulty"]["overall_level"] == "C2"


def test_model_config_can_switch_provider(client):
    updated = client.post(
        "/models/config",
        json={
            "provider": "mock",
            "ollama_model": "gemma4:test",
            "ollama_base_url": "http://localhost:11434",
            "mlx_model_path": "~/Models/mlx/gemma-4-e4b-it-OptiQ-4bit",
        },
    )
    assert updated.status_code == 200
    body = updated.json()
    assert body["provider"] == "mock"
    assert body["ollama_model"] == "gemma4:test"

    status = client.get("/models/status")
    assert status.status_code == 200
    assert status.json()["ollama_model"] == "gemma4:test"


def test_runtime_persisted_preset_applies_model_path(client):
    updated = client.post("/models/config", json={"preset_id": "gemma4-e2b-mlx"})
    assert updated.status_code == 200
    assert updated.json()["preset_id"] == "gemma4-e2b-mlx"
    assert "gemma-4-e2b" in updated.json()["mlx_model_path"]


def test_document_analysis_and_dictionary_flow(client):
    created = client.post(
        "/documents",
        json={"title": "Sleep study", "content": SAMPLE_TEXT, "source_type": "text"},
    )
    assert created.status_code == 200
    document = created.json()
    assert document["title"] == "Sleep study"

    listed = client.get("/documents")
    assert listed.status_code == 200
    assert listed.json()[0]["id"] == document["id"]

    analyzed = client.post(f"/documents/{document['id']}/analyze")
    assert analyzed.status_code == 200
    analysis = analyzed.json()
    assert analysis["document_id"] == document["id"]
    assert analysis["domain"]["primary_domain"]
    assert "quality_warnings" in analysis
    assert {term["term"] for term in analysis["terms"]} >= {
        "sleep deprivation",
        "cognitive performance",
        "longitudinal study",
    }
    assert all("learning_priority" in term for term in analysis["terms"])
    assert len(analysis["concepts"]) >= 1
    assert analysis["concepts"][0]["concept"]

    paper_map = client.get(f"/documents/{document['id']}/paper-map")
    assert paper_map.status_code == 200
    mapped_items = {
        item["text"].lower()
        for bucket in ("top_concepts", "top_terms")
        for item in paper_map.json()[bucket]
    }
    assert "training deep neural networks" not in mapped_items
    assert "inputs changes during training" not in mapped_items

    saved = client.post(
        "/dictionary/items",
        json={
            "item_type": "term",
            "text": "sleep deprivation",
            "meaning": "not getting enough sleep",
            "document_id": document["id"],
        },
    )
    assert saved.status_code == 200
    assert saved.json()["encounter_count"] == 1

    saved_again = client.post(
        "/dictionary/items",
        json={
            "item_type": "term",
            "text": "sleep deprivation",
            "meaning": "lack of sufficient sleep",
            "document_id": document["id"],
        },
    )
    assert saved_again.status_code == 200
    assert saved_again.json()["encounter_count"] == 2

    saved_concept = client.post(
        "/dictionary/items",
        json={
            "item_type": "concept",
            "text": analysis["concepts"][0]["concept"],
            "meaning": analysis["concepts"][0]["explanation"],
            "document_id": document["id"],
        },
    )
    assert saved_concept.status_code == 200

    items = client.get("/dictionary/items")
    assert items.status_code == 200
    assert len(items.json()) == 2
    term_item = next(item for item in items.json() if item["item_type"] == "term")
    assert term_item["view_count"] == 0

    viewed = client.post(f"/dictionary/items/{term_item['id']}/view")
    assert viewed.status_code == 200
    assert viewed.json()["view_count"] == 1
    assert viewed.json()["last_viewed_at"] is not None

    deleted = client.delete(f"/dictionary/items/{term_item['id']}")
    assert deleted.status_code == 204
    assert len(client.get("/dictionary/items").json()) == 1


def test_document_section_analysis_stays_on_parent_document(client):
    content = " ".join([SAMPLE_TEXT] * 14)
    created = client.post(
        "/documents",
        json={"title": "Long paper", "content": content, "source_type": "text"},
    )
    assert created.status_code == 200

    sections = client.get(f"/documents/{created.json()['id']}/sections")
    assert sections.status_code == 200
    assert len(sections.json()) >= 2
    assert sections.json()[1]["index"] == 1
    assert sections.json()[1]["section_number"] == 2
    assert sections.json()[1]["total_sections"] == len(sections.json())
    assert sections.json()[1]["char_count"] == len(sections.json()[1]["text"])
    assert sections.json()[1]["analyzed"] is False

    one_section = client.get(f"/documents/{created.json()['id']}/sections/1")
    assert one_section.status_code == 200
    assert one_section.json() == sections.json()[1]

    section = client.post(f"/documents/{created.json()['id']}/sections/1/analyze")
    assert section.status_code == 200
    body = section.json()
    assert body["document_id"] == created.json()["id"]
    assert any(warning.startswith("section:") for warning in body["quality_warnings"])
    assert body["terms"]
    assert body["summaries"]["one_line"]

    cached = client.post(f"/documents/{created.json()['id']}/sections/1/analyze")
    assert cached.status_code == 200
    assert cached.json() == body

    refreshed = client.post(f"/documents/{created.json()['id']}/sections/1/analyze?force=true")
    assert refreshed.status_code == 200
    assert refreshed.json()["document_id"] == created.json()["id"]
    assert refreshed.json()["summaries"]["one_line"]

    sections_after_analysis = client.get(f"/documents/{created.json()['id']}/sections")
    assert sections_after_analysis.status_code == 200
    assert sections_after_analysis.json()[1]["analyzed"] is True

    cached_section_analysis = client.get(f"/documents/{created.json()['id']}/sections/1/analysis")
    assert cached_section_analysis.status_code == 200
    assert cached_section_analysis.json()["document_id"] == created.json()["id"]
    assert cached_section_analysis.json()["summaries"]["one_line"]

    paper_map = client.get(f"/documents/{created.json()['id']}/paper-map")
    assert paper_map.status_code == 200
    assert paper_map.json()["total_sections"] >= 2
    assert paper_map.json()["analyzed_sections"] == [2]
    assert paper_map.json()["section_summaries"][0]["text"] == "Section 2"
    assert paper_map.json()["top_terms"]

    missing = client.post(f"/documents/{created.json()['id']}/sections/999/analyze")
    assert missing.status_code == 404

    deleted = client.delete(f"/documents/{created.json()['id']}")
    assert deleted.status_code == 204
    assert client.get(f"/documents/{created.json()['id']}/sections/1/analysis").status_code == 404


def test_cleanup_duplicate_documents_keeps_best_progress_record(client):
    content = " ".join([SAMPLE_TEXT] * 14)
    first = client.post("/documents", json={"title": "duplicate.pdf", "content": content, "source_type": "pdf"})
    second = client.post("/documents", json={"title": "duplicate.pdf", "content": content, "source_type": "pdf"})
    assert first.status_code == 200
    assert second.status_code == 200

    analyzed = client.post(f"/documents/{first.json()['id']}/sections/1/analyze")
    assert analyzed.status_code == 200

    cleanup = client.post("/documents/cleanup-duplicates")
    assert cleanup.status_code == 200
    body = cleanup.json()
    assert body["deleted_count"] == 1
    assert body["groups"][0]["kept_document_id"] == first.json()["id"]
    assert body["groups"][0]["deleted_document_ids"] == [second.json()["id"]]

    listed = client.get("/documents")
    assert listed.status_code == 200
    matching = [document for document in listed.json() if document["title"] == "duplicate.pdf"]
    assert [document["id"] for document in matching] == [first.json()["id"]]

    repeated = client.post("/documents/cleanup-duplicates")
    assert repeated.status_code == 200
    assert repeated.json()["deleted_count"] == 0


def test_staged_analysis_analyzes_next_unstudied_sections(client):
    content = " ".join([SAMPLE_TEXT] * 14)
    created = client.post(
        "/documents",
        json={"title": "Long staged paper", "content": content, "source_type": "text"},
    )
    assert created.status_code == 200

    staged = client.post(f"/documents/{created.json()['id']}/staged-analysis", json={"max_sections": 2})
    assert staged.status_code == 200
    body = staged.json()
    assert body["status"] == "completed"
    assert body["requested_sections"] == [1, 2]
    assert body["analyzed_sections"] == [1, 2]

    sections = client.get(f"/documents/{created.json()['id']}/sections")
    assert sections.status_code == 200
    assert sections.json()[0]["analyzed"] is True
    assert sections.json()[1]["analyzed"] is True

    staged_again = client.post(f"/documents/{created.json()['id']}/staged-analysis", json={"max_sections": 1})
    assert staged_again.status_code == 200
    assert staged_again.json()["requested_sections"] == [3]
    assert staged_again.json()["analyzed_sections"] == [3]


def test_page_batch_analysis_prepares_all_sections_on_requested_pdf_page(client):
    created = client.post(
        "/documents",
        json={
            "title": "Page batch PDF",
            "content": (
                "[[GEMMALENS_PDF_PAGE:1]]\n"
                "Abstract Batch Normalization reduces internal covariate shift. "
                "It uses mini-batch statistics to stabilize training.\n\n"
                "1 Introduction Stochastic gradient descent can be unstable in deep networks. "
                "The method allows higher learning rates.\n"
                "[[GEMMALENS_PDF_PAGE:2]]\n"
                "2 Background Attention mechanisms connect distant positions in a sequence. "
                "Self-attention improves parallelization."
            ),
            "source_type": "pdf",
        },
    )
    assert created.status_code == 200
    document_id = created.json()["id"]

    prepared = client.post(f"/documents/{document_id}/pages/1/analyze")
    assert prepared.status_code == 200
    body = prepared.json()
    assert body["requested_pages"] == [1]
    assert body["analyzed_pages"] == [1]
    assert body["analyzed_sections"] == [1, 2]

    sections = client.get(f"/documents/{document_id}/sections")
    assert sections.status_code == 200
    assert sections.json()[0]["analyzed"] is True
    assert sections.json()[1]["analyzed"] is True
    assert sections.json()[2]["analyzed"] is False

    first = client.get(f"/documents/{document_id}/sections/0/analysis")
    second = client.get(f"/documents/{document_id}/sections/1/analysis")
    assert first.status_code == 200
    assert second.status_code == 200
    assert first.json()["summaries"]["one_line"]
    assert second.json()["summaries"]["one_line"]


def test_page_batch_queue_skips_ready_pages_and_prepares_next_page(client):
    created = client.post(
        "/documents",
        json={
            "title": "Queued page batch PDF",
            "content": (
                "[[GEMMALENS_PDF_PAGE:1]]\n"
                "Abstract Batch Normalization reduces internal covariate shift. "
                "It uses mini-batch statistics to stabilize training.\n"
                "[[GEMMALENS_PDF_PAGE:2]]\n"
                "Background Attention mechanisms connect distant positions in a sequence. "
                "Self-attention improves parallelization."
            ),
            "source_type": "pdf",
        },
    )
    assert created.status_code == 200
    document_id = created.json()["id"]

    first = client.post(f"/documents/{document_id}/page-batches", json={"max_pages": 1})
    assert first.status_code == 200
    assert first.json()["requested_pages"] == [1]
    assert first.json()["analyzed_pages"] == [1]

    second = client.post(f"/documents/{document_id}/page-batches", json={"max_pages": 1})
    assert second.status_code == 200
    assert second.json()["requested_pages"] == [2]
    assert second.json()["analyzed_pages"] == [2]

    done = client.post(f"/documents/{document_id}/page-batches", json={"max_pages": 1})
    assert done.status_code == 200
    assert done.json()["status"] == "nothing_to_do"


def test_document_sections_include_pdf_page_labels_when_available(client):
    created = client.post(
        "/documents",
        json={
            "title": "Marked PDF",
            "content": (
                "[[GEMMALENS_PDF_PAGE:1]]\nAbstract\nBatch Normalization reduces internal covariate shift. "
                "[[GEMMALENS_PDF_PAGE:2]]\nThe method estimates mini-batch statistics."
            ),
            "source_type": "pdf",
        },
    )
    assert created.status_code == 200

    sections = client.get(f"/documents/{created.json()['id']}/sections")
    assert sections.status_code == 200
    assert [section["source_label"] for section in sections.json()] == ["PDF page 1", "PDF page 2"]


def test_upload_document_uses_ingestion_service(client):
    uploaded = client.post(
        "/documents/upload",
        files={"file": ("notes.md", b"# Methods\n\nWe analyze longitudinal study logs.", "text/markdown")},
    )
    assert uploaded.status_code == 200
    body = uploaded.json()
    assert body["title"] == "notes.md"
    assert body["source_type"] == "markdown"
    assert "longitudinal study logs" in body["content"]
    assert body["has_original_file"] is True

    original = client.get(f"/documents/{body['id']}/file")
    assert original.status_code == 200
    assert b"We analyze longitudinal study logs." in original.content

    empty = client.post(
        "/documents/upload",
        files={"file": ("empty.txt", b"\n\n", "text/plain")},
    )
    assert empty.status_code == 400


def test_upload_docx_document_uses_ingestion_service(client):
    from io import BytesIO

    from docx import Document

    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("Batch Normalization stabilizes deep network training.")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "term"
    doc.tables[0].rows[0].cells[1].text = "meaning"
    doc.save(buffer)

    uploaded = client.post(
        "/documents/upload",
        files={
            "file": (
                "notes.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert uploaded.status_code == 200
    body = uploaded.json()
    assert body["title"] == "notes.docx"
    assert body["source_type"] == "docx"
    assert body["has_original_file"] is True
    assert "Batch Normalization" in body["content"]
    assert "term | meaning" in body["content"]


def test_upload_legacy_doc_returns_actionable_error(client):
    uploaded = client.post(
        "/documents/upload",
        files={"file": ("legacy.doc", b"binary-doc-content", "application/msword")},
    )
    assert uploaded.status_code == 400
    assert ".doc files are not directly supported" in uploaded.json()["detail"]


def test_attach_original_file_to_existing_document(client):
    created = client.post(
        "/documents",
        json={"title": "Existing PDF text", "content": "Extracted text only", "source_type": "pdf"},
    )
    assert created.status_code == 200
    assert created.json()["has_original_file"] is False

    attached = client.post(
        f"/documents/{created.json()['id']}/file",
        files={"file": ("paper.pdf", b"%PDF-demo-bytes", "application/pdf")},
    )
    assert attached.status_code == 200
    assert attached.json()["has_original_file"] is True
    assert attached.json()["source_type"] == "pdf"

    original = client.get(f"/documents/{created.json()['id']}/file")
    assert original.status_code == 200
    assert original.content == b"%PDF-demo-bytes"


def test_missing_resources_return_404(client):
    assert client.get("/documents/missing").status_code == 404
    assert client.get("/documents/missing/analysis").status_code == 404
    assert client.delete("/dictionary/items/missing").status_code == 404
