from pathlib import Path
from uuid import uuid4

from fastapi import UploadFile

from app.core.config import get_settings
from app.schemas.document_schema import DocumentCreate


class DocumentIngestionError(ValueError):
    pass


class DocumentIngestionService:
    def __init__(self) -> None:
        self.settings = get_settings()

    def normalize_text(self, text: str) -> str:
        return "\n".join(line.strip() for line in text.replace("\r\n", "\n").splitlines() if line.strip())

    async def from_upload(self, file: UploadFile) -> DocumentCreate:
        name = file.filename or "Uploaded document"
        extension = name.rsplit(".", 1)[-1].lower() if "." in name else "txt"
        raw = await file.read()
        if extension == "pdf":
            content = self._extract_pdf(raw)
            source_type = "pdf"
        elif extension == "docx":
            content = self._extract_docx(raw)
            source_type = "docx"
        elif extension == "doc":
            raise DocumentIngestionError(
                "Legacy .doc files are not directly supported yet. Export it as .docx, .pdf with selectable text, or .txt and upload again."
            )
        elif extension in {"txt", "text", "md", "markdown"}:
            content = self._decode_text(raw)
            source_type = "markdown" if extension in {"md", "markdown"} else "text"
        else:
            raise DocumentIngestionError(
                f"Unsupported file type '.{extension}'. Supported uploads are PDF, DOCX, TXT, and Markdown."
            )
        normalized = self.normalize_text(content)
        if not normalized:
            if extension == "pdf":
                raise DocumentIngestionError(
                    "No extractable text found in this PDF. It may be scanned, image-only, encrypted, or need OCR. Try OCR first, upload a selectable-text PDF, or paste/upload the text."
                )
            if extension == "docx":
                raise DocumentIngestionError(
                    "No readable paragraphs found in this DOCX. It may contain only images, drawings, or unsupported embedded objects. Try exporting to PDF/TXT or paste the text."
                )
            raise DocumentIngestionError("No extractable text found in uploaded document")
        saved_path = self.save_original_file(raw, extension)
        return DocumentCreate(
            title=name,
            content=normalized,
            source_type=source_type,
            original_file_path=saved_path,
            original_mime_type=file.content_type,
        )

    def save_original_file(self, raw: bytes, extension: str) -> str | None:
        if not raw:
            return None
        storage_dir = Path(self.settings.upload_storage_dir)
        storage_dir.mkdir(parents=True, exist_ok=True)
        safe_extension = "".join(ch for ch in extension.lower() if ch.isalnum())[:12] or "bin"
        path = storage_dir / f"{uuid4()}.{safe_extension}"
        path.write_bytes(raw)
        return str(path)

    def _decode_text(self, raw: bytes) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr", "latin-1"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")

    def _extract_pdf(self, raw: bytes) -> str:
        try:
            from io import BytesIO

            from pypdf import PdfReader

            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:
                    return ""
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise DocumentIngestionError(f"Could not read PDF: {exc}") from exc

    def _extract_docx(self, raw: bytes) -> str:
        try:
            from io import BytesIO

            from docx import Document

            document = Document(BytesIO(raw))
            parts: list[str] = []
            parts.extend(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as exc:
            raise DocumentIngestionError(f"Could not read DOCX: {exc}") from exc
